import zipfile

from docx import Document

from ran2wiki.config import Config, PathsConfig
from ran2wiki.db import initialize
from ran2wiki.indexing import index_extracted
from ran2wiki.search import KnowledgeService


def make_chair(path, outcome: str) -> None:
    document = Document()
    document.add_paragraph("7.3 Dynamic UE capability")
    document.add_paragraph("R2-2501234 was discussed.")
    document.add_paragraph(outcome)
    document.save(path)


def test_final_chair_zip_is_selected_and_linked(tmp_path) -> None:
    config = Config(paths=PathsConfig(tmp_path, tmp_path / "wiki.db"))
    db = initialize(config.paths.database)
    draft = tmp_path / "ChairNotes_draft.docx"
    final_docx = tmp_path / "ChairNotes_final.docx"
    final_zip = tmp_path / "ChairNotes_final.zip"
    make_chair(draft, "FFS: Draft uncertainty")
    make_chair(final_docx, "Agreements: The final direction is approved.")
    with zipfile.ZipFile(final_zip, "w") as archive:
        archive.write(final_docx, final_docx.name)
    db.execute("INSERT INTO meetings(id,name,number,source_url) VALUES (1,'TSGR2_131bis',131,'https://meeting')")
    db.execute("INSERT INTO files(id,meeting_id,kind,filename,source_url,local_path) VALUES (1,1,'chair_notes','ChairNotes_draft.docx','https://draft',?)", (str(draft),))
    db.execute("INSERT INTO files(id,meeting_id,kind,filename,source_url,local_path) VALUES (2,1,'chair_notes','ChairNotes_final.zip','https://final',?)", (str(final_zip),))
    db.execute("INSERT INTO tdocs(id,tdoc_id,meeting_id,title) VALUES (1,'R2-2501234',1,'Dynamic UE capability')")
    db.commit()
    result = index_extracted(db, config)
    assert result["chair_sections"] == 1
    outcome = db.execute("SELECT outcome_type,text FROM chair_note_outcomes").fetchone()
    assert tuple(outcome) == ("AGREEMENT", "The final direction is approved.")
    assert db.execute("SELECT count(*) FROM chair_section_tdocs").fetchone()[0] == 1
    evidence = KnowledgeService(db).get_tdoc("R2-2501234")["chair_notes_references"][0]
    assert evidence["source_file"] == "ChairNotes_final.zip"


def test_zip_extraction_rejects_path_traversal(tmp_path) -> None:
    from ran2wiki.parsers.extract import extract_document

    archive_path = tmp_path / "bad.zip"
    with zipfile.ZipFile(archive_path, "w") as archive:
        archive.writestr("../escape.txt", "bad")
    try:
        extract_document(archive_path)
    except ValueError as exc:
        assert "unsafe ZIP member" in str(exc)
    else:
        raise AssertionError("unsafe archive was accepted")


def test_zip_extracts_nested_docx_text(tmp_path) -> None:
    from ran2wiki.parsers.extract import extract_document

    document_path = tmp_path / "R2-2501234.docx"
    make_chair(document_path, "Conclusion: nested evidence")
    archive_path = tmp_path / "R2-2501234.zip"
    with zipfile.ZipFile(archive_path, "w") as archive:
        archive.write(document_path, document_path.name)
    result = extract_document(archive_path)
    assert "nested evidence" in result["text"]
    assert result["blocks"][0]["locator"].startswith("R2-2501234.docx!")
