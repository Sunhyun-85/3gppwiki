from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

from ran2wiki.parsers.common import normalize_tdoc_id

ALIASES = {
    "tdoc_id": ("tdoc", "tdoc no", "tdoc number", "document number", "t-doc"),
    "title": ("title", "document title", "subject"),
    "source_company": ("source", "company", "source/company", "author"),
    "agenda": ("agenda", "agenda item", "ai"),
    "work_item": ("wi/si", "wis", "work item", "work item / study item", "feature"),
    "status": ("status", "document status", "result"),
    "revision": ("revision", "revised from", "revision of", "rev of"),
    "comments": ("comments", "comment", "notes"),
    "filename": ("file name", "filename", "file"),
}


def _header(value: object) -> str:
    return re.sub(r"\s+", " ", str(value or "").strip().lower())


def normalize_columns(headers: Iterable[object]) -> dict[int, str]:
    result: dict[int, str] = {}
    for index, raw in enumerate(headers):
        value = _header(raw)
        for canonical, aliases in ALIASES.items():
            if value in aliases:
                result[index] = canonical
                break
    return result


@dataclass(frozen=True)
class TDocRow:
    tdoc_id: str
    title: str | None = None
    source_company: str | None = None
    agenda: str | None = None
    work_item: str | None = None
    status: str | None = None
    revision: str | None = None
    comments: str | None = None
    filename: str | None = None


def parse_rows(rows: Iterable[Iterable[object]]) -> list[TDocRow]:
    materialized = [list(row) for row in rows]
    header_at = next((i for i, row in enumerate(materialized[:30]) if "tdoc_id" in normalize_columns(row).values()), None)
    if header_at is None:
        raise ValueError("could not locate a TDoc number header in the first 30 rows")
    columns = normalize_columns(materialized[header_at])
    parsed: list[TDocRow] = []
    for row in materialized[header_at + 1:]:
        values = {canonical: (str(row[i]).strip() if i < len(row) and row[i] is not None else None) for i, canonical in columns.items()}
        tdoc_id = normalize_tdoc_id(values.get("tdoc_id"))
        if tdoc_id:
            values["tdoc_id"] = tdoc_id
            parsed.append(TDocRow(**values))
    return parsed


def parse_workbook(path: str | Path) -> list[TDocRow]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    best: list[TDocRow] = []
    for sheet in workbook.worksheets:
        try:
            candidates = parse_rows(sheet.iter_rows(values_only=True))
        except ValueError:
            continue
        if len(candidates) > len(best):
            best = candidates
    if not best:
        raise ValueError(f"no TDoc rows found in {path}")
    return best

