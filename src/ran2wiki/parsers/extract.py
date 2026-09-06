from __future__ import annotations

import json
import zipfile
import tempfile
from pathlib import Path
from typing import Any


MAX_ARCHIVE_MEMBER = 100 * 1024 * 1024
MAX_ARCHIVE_TOTAL = 500 * 1024 * 1024


def _safe_members(archive: zipfile.ZipFile) -> list[zipfile.ZipInfo]:
    members = []
    total = 0
    for item in archive.infolist():
        candidate = Path(item.filename)
        if item.is_dir():
            continue
        if candidate.is_absolute() or ".." in candidate.parts:
            raise ValueError(f"unsafe ZIP member path: {item.filename}")
        if item.file_size > MAX_ARCHIVE_MEMBER:
            raise ValueError(f"ZIP member is too large: {item.filename}")
        total += item.file_size
        if total > MAX_ARCHIVE_TOTAL:
            raise ValueError("ZIP expanded size exceeds safety limit")
        members.append(item)
    return members


def extract_document(path: str | Path, *, _depth: int = 0) -> dict[str, Any]:
    source = Path(path)
    suffix = source.suffix.lower()
    blocks: list[dict[str, Any]] = []
    if suffix == ".docx":
        from ran2wiki.parsers.chair_notes import docx_blocks
        blocks = [{"type": "block", "ordinal": i, "text": text} for i, text in enumerate(docx_blocks(source), 1)]
    elif suffix == ".pptx":
        from pptx import Presentation
        for number, slide in enumerate(Presentation(source).slides, 1):
            text = "\n".join(shape.text for shape in slide.shapes if hasattr(shape, "text") and shape.text)
            blocks.append({"type": "slide", "ordinal": number, "locator": f"slide:{number}", "text": text})
    elif suffix in {".xlsx", ".xlsm"}:
        from openpyxl import load_workbook
        workbook = load_workbook(source, read_only=True, data_only=True)
        for sheet in workbook.worksheets:
            for number, row in enumerate(sheet.iter_rows(values_only=True), 1):
                text = " | ".join(str(value) for value in row if value is not None)
                if text:
                    blocks.append({"type": "row", "ordinal": len(blocks) + 1, "locator": f"{sheet.title}!{number}", "text": text})
    elif suffix == ".pdf":
        from pypdf import PdfReader
        for number, page in enumerate(PdfReader(source).pages, 1):
            blocks.append({"type": "page", "ordinal": number, "locator": f"page:{number}", "text": page.extract_text() or ""})
    elif suffix in {".txt", ".md", ".csv"}:
        blocks = [{"type": "text", "ordinal": 1, "text": source.read_text(encoding="utf-8", errors="replace")}]
    elif suffix == ".zip":
        with zipfile.ZipFile(source) as archive:
            members = _safe_members(archive)
            with tempfile.TemporaryDirectory(prefix="ran2wiki-extract-") as temporary:
                root = Path(temporary)
                for item in members:
                    member_path = root / item.filename
                    member_path.parent.mkdir(parents=True, exist_ok=True)
                    with archive.open(item) as incoming, member_path.open("wb") as outgoing:
                        while chunk := incoming.read(1024 * 1024):
                            outgoing.write(chunk)
                    if member_path.suffix.lower() in {".docx", ".pptx", ".xlsx", ".xlsm", ".pdf", ".txt"}:
                        nested = extract_document(member_path, _depth=_depth + 1)
                        for nested_block in nested["blocks"]:
                            block = dict(nested_block)
                            locator = block.get("locator", f"block:{block.get('ordinal', 0)}")
                            block["locator"] = f"{item.filename}!{locator}"
                            block["member"] = item.filename
                            block["ordinal"] = len(blocks) + 1
                            blocks.append(block)
                    else:
                        blocks.append({"type": "archive_member", "ordinal": len(blocks) + 1,
                                       "locator": item.filename, "text": item.filename, "size": item.file_size})
    else:
        raise ValueError(f"unsupported document type: {suffix}")
    return {"source": source.name, "format": suffix.removeprefix("."), "blocks": blocks, "text": "\n".join(block.get("text", "") for block in blocks)}


def write_extracted(result: dict[str, Any], destination: str | Path) -> None:
    output = Path(destination)
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
